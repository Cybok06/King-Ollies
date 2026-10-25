"""Private, in-memory application exports. Never fetch remote assets or embed HTML."""
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape


def date_text(value):
    if not isinstance(value, datetime): return 'Not recorded'
    if value.tzinfo: value = value.astimezone(timezone.utc)
    return value.strftime('%d %b %Y, %H:%M GMT')


def summary(application, customer):
    return [
        ('Application reference', str(application['_id'])),
        ('Service', application.get('service_name', 'Service application')),
        ('Customer', customer.get('username') or customer.get('name') or 'Customer'),
        ('Email', customer.get('email') or 'Not provided'),
        ('Phone', customer.get('phone') or 'Not provided'),
        ('Submitted', date_text(application.get('created_at'))),
        ('Application progress', application.get('progress_status', application.get('status', 'submitted')).replace('_', ' ').title()),
        ('Payment status', str(application.get('status', 'submitted')).replace('_', ' ').title()),
        ('Invoice amount (GHS)', application.get('quote_amount') if application.get('quote_amount') is not None else 'Not issued'),
        ('Paid amount (GHS)', application.get('paid_amount') if application.get('paid_at') else 'Not confirmed'),
        ('Paid on', date_text(application.get('paid_at'))),
        ('Payment method', application.get('payment_method') or 'Not recorded'),
    ]


def export_excel(application, customer):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, headers, rows in [
        ('Application summary', ['Field', 'Value'], summary(application, customer)),
        ('Submitted details', ['Section', 'Field', 'Submitted value'], [
            [answer.get('section') or 'General', answer.get('label', ''), answer.get('value', '')]
            for answer in application.get('answers', [])]),
    ]:
        sheet = workbook.create_sheet(title)
        sheet.append(["King Ollies | " + title])
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        sheet.cell(1, 1).font = Font(name='Calibri', size=17, bold=True, color='FFFFFF')
        sheet.cell(1, 1).fill = PatternFill('solid', fgColor='1A1A1E')
        sheet.row_dimensions[1].height = 36
        sheet.append(['Private application record · ' + str(application['_id'])])
        sheet.append(headers)
        for row in rows:
            sheet.append([value if isinstance(value, (int, float)) else str(value or '') for value in row])
            for cell in sheet[sheet.max_row]:
                # Treat all user-provided strings as text, including leading '='.
                if isinstance(cell.value, str): cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.fill = PatternFill('solid', fgColor='F8F7F3' if sheet.max_row % 2 == 0 else 'FFFFFF')
                cell.font = Font(name='Calibri', size=11, color='1A1A1E')
            longest = max((len(str(value)) for value in row), default=1)
            sheet.row_dimensions[sheet.max_row].height = min(300, max(28, (longest // 65 + 1) * 16))
        for cell in sheet[3]:
            cell.fill = PatternFill('solid', fgColor='D99A24')
            cell.font = Font(bold=True, color='0B0B0D')
        for index in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 80 if index == len(headers) else 30
        sheet.freeze_panes = 'A4'
        sheet.auto_filter.ref = f'A3:{get_column_letter(len(headers))}{sheet.max_row}'
        sheet.sheet_view.showGridLines = False
        sheet.print_title_rows = '1:3'
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def export_pdf(application, customer):
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    font_dir = Path(reportlab.__file__).parent / 'fonts'
    if 'ApplicationSans' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('ApplicationSans', str(font_dir / 'Vera.ttf')))
        pdfmetrics.registerFont(TTFont('ApplicationSansBold', str(font_dir / 'VeraBd.ttf')))
    regular = ParagraphStyle('body', fontName='ApplicationSans', fontSize=10, leading=15, textColor=colors.HexColor('#25252B'), spaceAfter=8, splitLongWords=True)
    label = ParagraphStyle('label', parent=regular, fontName='ApplicationSansBold', fontSize=9, textColor=colors.HexColor('#55555F'), spaceAfter=4, keepWithNext=True)
    heading = ParagraphStyle('section', parent=regular, fontName='ApplicationSansBold', fontSize=14, leading=20, spaceBefore=18, spaceAfter=10, keepWithNext=True)
    title = ParagraphStyle('title', parent=heading, fontSize=22, leading=28, spaceBefore=0)
    def paragraph(value, style=regular):
        return Paragraph(escape(str(value or '—')).replace('\n', '<br/>'), style)
    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=A4, rightMargin=44, leftMargin=44, topMargin=110, bottomMargin=60,
                            title='Application - '+str(application.get('service_name', '')), author="King Ollies")
    generated = date_text(datetime.now(timezone.utc))
    def page_frame(canvas, document):
        width, height = A4
        canvas.saveState()
        canvas.setFillColor(colors.HexColor('#0B0B0D'))
        canvas.rect(0, height-82, width, 82, fill=1, stroke=0)
        canvas.setFillColor(colors.HexColor('#D99A24'))
        canvas.rect(0, height-85, width, 3, fill=1, stroke=0)
        canvas.setFont('ApplicationSansBold', 20)
        canvas.drawString(44, height-39, "King Ollies")
        canvas.setFont('ApplicationSans', 9)
        canvas.setFillColor(colors.white)
        canvas.drawString(44, height-60, 'AGENCY AND DIGITAL SERVICES  /  APPLICATION RECORD')
        canvas.setStrokeColor(colors.HexColor('#DDDDDF'))
        canvas.line(44, 44, width-44, 44)
        canvas.setFont('ApplicationSans', 8)
        canvas.setFillColor(colors.HexColor('#66666F'))
        canvas.drawString(44, 29, 'Confidential · Generated '+generated)
        canvas.drawRightString(width-44, 29, f'Page {document.page}')
        canvas.restoreState()
    story = [paragraph(application.get('service_name', 'Service application'), title), paragraph('Submitted application details')]
    entries = summary(application, customer)
    rows = [[paragraph(entries[i][0], label), paragraph(entries[i][1]),
             paragraph(entries[i+1][0], label), paragraph(entries[i+1][1])]
            for i in range(0, len(entries), 2)]
    table = Table(rows, colWidths=[96, (A4[0]-280)/2, 96, (A4[0]-280)/2], hAlign='LEFT')
    table.setStyle(TableStyle([
        ('VALIGN',(0,0),(-1,-1),'TOP'), ('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#F8F7F3')),
        ('LINEBELOW',(0,0),(-1,-1),.4,colors.HexColor('#E6E4DE')),
        ('LEFTPADDING',(0,0),(-1,-1),12), ('RIGHTPADDING',(0,0),(-1,-1),12),
        ('TOPPADDING',(0,0),(-1,-1),8), ('BOTTOMPADDING',(0,0),(-1,-1),6)]))
    story.extend([table, Spacer(1, 8)])
    current_section = None
    for answer in application.get('answers', []):
        section = answer.get('section') or 'Submitted details'
        if section != current_section:
            story.append(paragraph(section, heading))
            current_section = section
        story.extend([paragraph(answer.get('label', ''), label), paragraph(answer.get('value', '')), Spacer(1, 5)])
    if application.get('notes'):
        story.extend([paragraph('Service notes', heading), paragraph(application['notes'])])
    doc.build(story, onFirstPage=page_frame, onLaterPages=page_frame)
    stream.seek(0)
    return stream
